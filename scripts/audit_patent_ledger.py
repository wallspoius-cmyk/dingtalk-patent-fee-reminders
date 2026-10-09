"""Read-only patent ledger audit. Requires openpyxl; never saves the workbook."""
import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ALIASES = {
    'name': ['专利名称', '名称'],
    'application_number': ['申请号', '专利号'],
    'due_date': ['缴费日', '缴费截止日', '缴费日期'],
    'amount': ['金额（元）', '金额', '年费金额'],
    'payment_status': ['缴费状态'],
    'application_date': ['申请日', '申请日期'],
    'fee_name': ['年费名称'],
    'source_category': ['原表分类', '专利类型'],
    'source_note': ['原表备注', '备注'],
    'source_sequence': ['原表序号', '序号'],
}


def scalar(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def iso_date(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, str):
        match = re.fullmatch(r'\s*(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?\s*', value)
        if match:
            try:
                return date(*map(int, match.groups())).isoformat()
            except ValueError:
                pass
    return None


def displayed_number(cell):
    value = cell.value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Known patent-ledger formats such as 0.0_ preserve a displayed .0.
        fmt = cell.number_format.split(';')[0]
        if re.fullmatch(r'0\.0(?:_.)?\s*', fmt):
            return format(Decimal(str(value)).quantize(Decimal('0.1')), 'f')
        return str(value)
    return None if value is None else str(value)


def audit(path, sheet_name=None, header_row=1):
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        if sheet_name is None and len(wb.sheetnames) != 1:
            raise ValueError('Workbook has multiple sheets; select one with --sheet.')
        sheet = wb[sheet_name] if sheet_name else wb.active
        headers = {}
        for cell in sheet[header_row]:
            if cell.value is not None:
                headers.setdefault(str(cell.value).strip(), []).append(cell.column)
        mapping = {}
        warnings = []
        for key, aliases in ALIASES.items():
            candidates = [(alias, col) for alias in aliases for col in headers.get(alias, [])]
            if len(candidates) > 1:
                warnings.append(f'{key}: multiple candidate columns; selected {candidates[0][0]}. Review mapping.')
            if candidates:
                mapping[key] = candidates[0][1]
        if 'name' not in mapping or 'application_number' not in mapping:
            raise ValueError('Cannot identify patent name and application number headers.')
        records = []
        for row_num in range(header_row + 1, sheet.max_row + 1):
            cells = {key: sheet.cell(row_num, col) for key, col in mapping.items()}
            if all(cells[key].value is None for key in ('name', 'application_number')):
                continue
            item = {'source_row': row_num, 'fields': {}, 'issues': []}
            for key, cell in cells.items():
                item['fields'][key] = {
                    'cell': cell.coordinate, 'value': scalar(cell.value),
                    'data_type': cell.data_type, 'number_format': cell.number_format,
                }
            id_cell = cells['application_number']
            number = displayed_number(id_cell)
            item['display_application_number'] = number
            if id_cell.data_type == 'n' and id_cell.value is not None:
                item['issues'].append('Application number stored as numeric; preserve displayed format and verify.')
            if number and not re.fullmatch(r'\d{12}\.[0-9Xx]', number):
                item['issues'].append('Application number format needs review; original retained.')
            due_value = cells.get('due_date').value if 'due_date' in cells else None
            item['parsed_due_date'] = iso_date(due_value)
            if due_value is None:
                item['issues'].append('Missing due date; exclude from active reminders until confirmed.')
            elif item['parsed_due_date'] is None:
                item['issues'].append('Due date not parsed unambiguously; review original value.')
            status = cells.get('payment_status').value if 'payment_status' in cells else None
            item['explicit_payment_status'] = scalar(status)
            if status is None:
                item['issues'].append('Payment status unconfirmed; do not infer from authorization notes.')
            records.append(item)
        counts = Counter(r['display_application_number'] for r in records if r['display_application_number'])
        report = {
            'source_sha256': before, 'sheet': sheet.title,
            'merged_ranges': [str(r) for r in sheet.merged_cells.ranges],
            'column_mapping': {key: sheet.cell(header_row, col).coordinate for key, col in mapping.items()},
            'mapping_warnings': warnings, 'record_count': len(records),
            'dated_record_count': sum(r['parsed_due_date'] is not None for r in records),
            'unconfirmed_status_count': sum(r['explicit_payment_status'] is None for r in records),
            'duplicate_application_numbers': [key for key, count in counts.items() if count > 1],
            'records': records,
        }
    finally:
        wb.close()
    if hashlib.sha256(path.read_bytes()).hexdigest() != before:
        raise RuntimeError('Source changed during audit; re-read before using this report.')
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('--out', required=True, type=Path)
    parser.add_argument('--sheet')
    parser.add_argument('--header-row', type=int, default=1)
    args = parser.parse_args()
    if args.header_row < 1:
        parser.error('--header-row must be positive.')
    if args.input.resolve() == args.out.resolve():
        parser.error('Output cannot overwrite the source workbook.')
    report = audit(args.input, args.sheet, args.header_row)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({key: report[key] for key in ('record_count', 'dated_record_count', 'unconfirmed_status_count')}, ensure_ascii=False))


if __name__ == '__main__':
    main()

